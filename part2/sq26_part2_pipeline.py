"""
sq26_part2_pipeline.py

Seeding QDArchive - Part 2: Classification - single-file pipeline.

Combines everything into one script:
  Step 1  Schema migration          (add type / primary_class / secondary_class columns)
  Step 1  Project-type classifier   (QDA_PROJECT / QD_PROJECT / OTHER_PROJECT / NOT_A_PROJECT)
  Step 2+3 ISIC Rev.5 classifier    (TF-IDF match against ISIC section+division taxonomy)
  Step 4  Report generation         (XLSX extract + PDF report per repository)

Usage
-----
    python3 sq26_part2_pipeline.py \
        --db 23129103-seeding.db \
        --isic-xlsx ISIC5_Exp_Notes_11Mar2024.xlsx \
        --out-prefix 23129103-sq26-classification

This will produce, next to the input DB:
    23129103-sq26-classification.db            (copy of input DB + new columns/values)
    23129103-sq26-classification-results.xlsx  (Step 4c deliverable)
    23129103-sq26-classification-report.pdf    (Step 4d deliverable)
    isic_divisions.json                        (cached ISIC taxonomy extract, reused on reruns)

Method notes / assumptions are documented inline at each step.
"""
import argparse
import json
import os
import shutil
import sqlite3
import textwrap
from collections import Counter, defaultdict

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


# ===========================================================================
# STEP 1a - Schema migration
# ===========================================================================

def column_exists(cur, table, column):
    cur.execute(f"PRAGMA table_info({table})")
    return any(row[1] == column for row in cur.fetchall())


def migrate_schema(conn):
    """Adds the Part 2 columns to PROJECTS / FILES, without touching data."""
    cur = conn.cursor()
    migrations = [
        ("PROJECTS", "type", "TEXT"),
        ("PROJECTS", "primary_class", "TEXT"),
        ("PROJECTS", "secondary_class", "TEXT"),
        ("FILES", "primary_class", "TEXT"),
    ]
    for table, column, coltype in migrations:
        if column_exists(cur, table, column):
            print(f"[schema] skip   {table}.{column} already exists")
            continue
        cur.execute(f"ALTER TABLE {table} ADD COLUMN {column} {coltype}")
        print(f"[schema] added  {table}.{column} {coltype}")
    conn.commit()


# ===========================================================================
# STEP 1b - Project-type classifier
# ===========================================================================
# Extension lists sourced from QDA_File_Extensions_Formats.xlsx (Standard/
# REFI-QDA, MaxQDA, NVivo, ATLAS.ti, QDAcity, QDA Miner, f4analyse, Quirkos)
# plus what actually appears in this student's FILES table (qdpx, qdc,
# atlproj, atlproj22, atlproj9, hpr7, mx20, ...).

QDA_EXTENSIONS = {
    "qdpx", "qdc",                                              # REFI-QDA standard
    "mqda", "mqbac", "mqtc", "mqex", "mqmtr",                   # MaxQDA
    "mx24", "mx24bac", "mc24", "mex24", "mx22", "mx20", "mx18",
    "mx12", "mx11", "mx5", "mx4", "mx3", "mx2", "m2k",
    "loa", "sea", "mtr", "mod", "mex22",
    "nvp", "nvpx",                                               # NVivo
    "atlasproj", "atlproj", "atlproj22", "atlproj9", "hpr7",     # ATLAS.ti
    "ppj", "pprj", "qlt",                                        # QDA Miner
    "f4p",                                                       # f4analyse
    "qpd",                                                       # Quirkos
}

# Primary data files: the qualitative input data itself (interview
# transcripts, articles, audio/video recordings, images of documents, etc.)
PRIMARY_DATA_EXTENSIONS = {
    "pdf", "doc", "docx", "txt", "rtf", "odt", "wpd", "wp5",
    "htm", "html", "md", "rst", "tex",
    "jpg", "jpeg", "png", "tif", "tiff", "bmp", "gif",
    "wav", "mp3", "mp4", "mov", "wma", "m4a", "flac", "avi", "mkv",
    "m4v", "3gp", "3gpp", "mpg", "mka", "mks", "mxf",
    "vtt", "srt",
}

# Other recognisable data files (spreadsheets, stats/GIS datasets,
# structured data, code, archives) -> OTHER_PROJECT if nothing above applies.
OTHER_DATA_EXTENSIONS = {
    "csv", "tsv", "tab", "xls", "xlsx", "xlsm", "ods", "sav", "por", "sps",
    "dta", "rdata", "rda", "rhistory", "rmd", "rproj", "jasp", "amosoutput",
    "json", "xml", "cmdi", "geojson", "kml", "kmz", "gpx",
    "shp", "shx", "dbf", "prj", "sbn", "sbx", "gdbtable", "gdbtablx",
    "gdbindexes", "mxd", "qml", "gpkg", "tpk", "lyr",
    "zip", "rar", "7z", "tar", "gz", "bz", "bz2",
    "sql", "sqlite", "db", "mdb", "accdb",
    "eeg", "edf", "vhdr", "vmrk",
    "ipynb", "py", "r", "m",
}


def classify_project_type_from_extensions(file_types):
    exts = {(ft or "").strip().lower().lstrip(".") for ft in file_types}
    if exts & QDA_EXTENSIONS:
        return "QDA_PROJECT"
    if exts & PRIMARY_DATA_EXTENSIONS:
        return "QD_PROJECT"
    if exts & OTHER_DATA_EXTENSIONS:
        return "OTHER_PROJECT"
    return "NOT_A_PROJECT"


def classify_project_types(conn):
    cur = conn.cursor()
    cur.execute("SELECT id FROM PROJECTS")
    project_ids = [r[0] for r in cur.fetchall()]

    cur.execute("SELECT project_id, file_type FROM FILES")
    files_by_project = defaultdict(list)
    for project_id, file_type in cur.fetchall():
        files_by_project[project_id].append(file_type)

    counts = Counter()
    for pid in project_ids:
        ptype = classify_project_type_from_extensions(files_by_project.get(pid, []))
        counts[ptype] += 1
        cur.execute("UPDATE PROJECTS SET type = ? WHERE id = ?", (ptype, pid))
    conn.commit()

    print("[type] Project type distribution:")
    for ptype in ("QDA_PROJECT", "QD_PROJECT", "OTHER_PROJECT", "NOT_A_PROJECT"):
        print(f"  {ptype:15s} {counts[ptype]:6d}")
    print(f"  {'TOTAL':15s} {sum(counts.values()):6d}")
    return counts


# ===========================================================================
# STEP 2+3 - ISIC Rev.5 classifier
# ===========================================================================
# We don't have the actual downloaded file *contents* available - only the
# metadata database - so this is a metadata-driven content classifier:
#   1. Build one TF-IDF reference document per ISIC division (title + intro
#      + includes + includes-also, aggregated up from groups/classes).
#   2. Build one TF-IDF query document per project (title + description +
#      keywords).
#   3. Vectorize both with a shared vocabulary, rank divisions by cosine
#      similarity per project.
#   4. primary_class = best match; secondary_class = second-best, but only
#      if it scores >= 85% of the best (otherwise it's noise, not a genuine
#      second candidate).
# Files inherit their parent project's primary_class (documented limitation:
# no file-content classification is possible without the actual data files).

SECONDARY_CLASS_RATIO = 0.85


def extract_isic_divisions(isic_xlsx_path, cache_path="isic_divisions.json"):
    if os.path.exists(cache_path):
        with open(cache_path) as f:
            return json.load(f)["divisions"]

    import openpyxl
    import re

    wb = openpyxl.load_workbook(isic_xlsx_path, data_only=True)
    ws = wb["ISIC5"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    sections, divisions = {}, {}
    current_division = None
    for row in rows:
        code_full = row[0]
        if not code_full:
            continue
        code_full = str(code_full).strip()
        title, intro = row[2] or "", row[3] or ""
        includes, includes_also = row[4] or "", row[5] or ""
        text_piece = " ".join([str(title), str(intro), str(includes), str(includes_also)])

        if re.fullmatch(r"[A-Z]", code_full):
            sections[code_full] = {"title": title, "text": text_piece}
            current_division = None
        elif re.fullmatch(r"[A-Z]\d{2}", code_full):
            divisions[code_full] = {"title": title, "text": text_piece, "section": code_full[0]}
            current_division = code_full
        else:
            if current_division:
                divisions[current_division]["text"] += " " + text_piece

    with open(cache_path, "w") as f:
        json.dump({"sections": sections, "divisions": divisions}, f)
    print(f"[isic]  extracted {len(divisions)} divisions -> cached to {cache_path}")
    return divisions


def classify_isic(conn, divisions):
    codes = list(divisions.keys())
    div_texts = [divisions[c]["text"] for c in codes]

    cur = conn.cursor()
    cur.execute(
        "SELECT id, title, description FROM PROJECTS WHERE type IN ('QDA_PROJECT', 'QD_PROJECT')"
    )
    projects = cur.fetchall()

    cur.execute("SELECT project_id, keyword FROM KEYWORDS")
    keywords_by_project = defaultdict(list)
    for pid, kw in cur.fetchall():
        keywords_by_project[pid].append(kw or "")

    project_ids, project_texts = [], []
    for pid, title, description in projects:
        kw_text = " ".join(keywords_by_project.get(pid, []))
        project_ids.append(pid)
        project_texts.append(" ".join([title or "", description or "", kw_text]))

    print(f"[isic]  classifying {len(project_ids)} projects against {len(codes)} ISIC divisions ...")

    vectorizer = TfidfVectorizer(stop_words="english", max_features=20000, ngram_range=(1, 2), min_df=1)
    combined = div_texts + project_texts
    tfidf = vectorizer.fit_transform(combined)
    div_vecs = tfidf[: len(div_texts)]
    proj_vecs = tfidf[len(div_texts):]
    sims = cosine_similarity(proj_vecs, div_vecs)

    unmatched = 0
    for i, pid in enumerate(project_ids):
        row = sims[i]
        order = row.argsort()[::-1]
        best_idx = order[0]
        best_score = row[best_idx]

        if best_score <= 0:
            primary_class, secondary_class = None, None
            unmatched += 1
        else:
            primary_class = codes[best_idx]
            second_idx = order[1]
            second_score = row[second_idx]
            secondary_class = (
                codes[second_idx]
                if second_score > 0 and second_score >= SECONDARY_CLASS_RATIO * best_score
                else None
            )

        cur.execute(
            "UPDATE PROJECTS SET primary_class = ?, secondary_class = ? WHERE id = ?",
            (primary_class, secondary_class, pid),
        )

    # Propagate to FILES.
    cur.execute(
        """
        UPDATE FILES
        SET primary_class = (SELECT primary_class FROM PROJECTS WHERE PROJECTS.id = FILES.project_id)
        WHERE project_id IN (SELECT id FROM PROJECTS WHERE type IN ('QDA_PROJECT', 'QD_PROJECT'))
        """
    )
    conn.commit()
    print(f"[isic]  done. {unmatched} projects had no vocabulary overlap (left NULL).")


# ===========================================================================
# STEP 4c - XLSX export
# ===========================================================================

def export_xlsx(conn, out_path):
    df = pd.read_sql_query(
        """
        SELECT
            p.repository_id  AS repository_id,
            p.type           AS project_type,
            p.title          AS project_title,
            p.primary_class  AS primary_class,
            p.secondary_class AS secondary_class,
            (SELECT COUNT(*) FROM FILES f WHERE f.project_id = p.id) AS no_project_files
        FROM PROJECTS p
        ORDER BY p.repository_id, p.id
        """,
        conn,
    )
    df.to_excel(out_path, index=False, sheet_name="results")

    wb = load_workbook(out_path)
    ws = wb["results"]
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    body_font = Font(name="Arial")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
    for i, w in enumerate([14, 16, 45, 14, 15, 16], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    wb.save(out_path)
    print(f"[xlsx]  {len(df)} rows -> {out_path}")


# ===========================================================================
# STEP 4d - PDF report
# ===========================================================================

def full_class_name(code, div_titles):
    if code is None:
        return "Unclassified"
    t = div_titles.get(code, {}).get("title", "?").strip()
    return f"{code} \u2013 {t}"


def get_repo_stats(conn, repo_id):
    cur = conn.cursor()
    cur.execute("SELECT type, COUNT(*) FROM PROJECTS WHERE repository_id = ? GROUP BY type", (repo_id,))
    type_counts = dict(cur.fetchall())
    cur.execute(
        """
        SELECT primary_class, COUNT(*) FROM PROJECTS
        WHERE repository_id = ? AND type IN ('QDA_PROJECT','QD_PROJECT') AND primary_class IS NOT NULL
        GROUP BY primary_class ORDER BY COUNT(*) DESC
        """,
        (repo_id,),
    )
    class_counts = cur.fetchall()
    return type_counts, class_counts


def generate_report(conn, div_titles, repos, comments_by_repo, out_path):
    def title_page(pdf):
        fig = plt.figure(figsize=(11, 8.5))
        fig.text(0.5, 0.62, "Seeding QDArchive", ha="center", fontsize=26, fontweight="bold")
        fig.text(0.5, 0.55, "Part 2: Classification Results Report", ha="center", fontsize=16)
        fig.text(0.5, 0.45, "Student ID: 23129103", ha="center", fontsize=12)
        fig.text(0.5, 0.41, "SQ26 \u2013 Prof. Dr. Dirk Riehle, FAU Erlangen-N\u00fcrnberg", ha="center", fontsize=11)
        repo_line = ", ".join(f"#{rid} {rname}" for rid, rname in repos)
        fig.text(0.5, 0.37, f"Repositories: {repo_line}", ha="center", fontsize=11)
        plt.axis("off")
        pdf.savefig(fig)
        plt.close(fig)

    def type_distribution_page(pdf, repo_name, type_counts):
        order = ["QDA_PROJECT", "QD_PROJECT", "OTHER_PROJECT", "NOT_A_PROJECT"]
        values = [type_counts.get(t, 0) for t in order]
        fig, ax = plt.subplots(figsize=(11, 6))
        bars = ax.bar(order, values, color=["#2E7D32", "#1565C0", "#F9A825", "#B71C1C"])
        for b, v in zip(bars, values):
            ax.text(b.get_x() + b.get_width() / 2, b.get_height(), str(v), ha="center", va="bottom", fontsize=11)
        ax.set_title(f"Repository: {repo_name} \u2014 Project type distribution", fontsize=14)
        ax.set_ylabel("Number of projects")
        plt.tight_layout()
        pdf.savefig(fig)
        plt.close(fig)

    def histogram_page(pdf, repo_name, class_counts):
        if not class_counts:
            fig = plt.figure(figsize=(11, 6))
            fig.text(0.5, 0.55, f"Repository: {repo_name}", ha="center", fontsize=16, fontweight="bold")
            fig.text(0.5, 0.45,
                     "No classifiable projects (QDA_PROJECT / QD_PROJECT) with a primary class\n"
                     "were found for this repository \u2014 see comments.",
                     ha="center", fontsize=12)
            plt.axis("off")
            pdf.savefig(fig)
            plt.close(fig)
            return
        top20 = class_counts[:20]
        labels = [full_class_name(c, div_titles) for c, _ in top20]
        values = [v for _, v in top20]
        fig, ax = plt.subplots(figsize=(11, 9))
        y_pos = range(len(labels))
        bars = ax.barh(y_pos, values, color="#1565C0")
        ax.set_yticks(y_pos)
        ax.set_yticklabels(["\n".join(textwrap.wrap(l, 55)) for l in labels], fontsize=8)
        ax.invert_yaxis()
        ax.set_xlabel("Number of projects")
        ax.set_title(f"Repository: {repo_name} \u2014 Top 20 primary ISIC classes", fontsize=13)
        for b, v in zip(bars, values):
            ax.text(b.get_width(), b.get_y() + b.get_height() / 2, f" {v}", va="center", fontsize=8)
        plt.tight_layout()
        pdf.savefig(fig)
        plt.close(fig)

    def table_page(pdf, repo_name, class_counts):
        if not class_counts:
            return
        top20 = class_counts[:20]
        total = sum(v for _, v in class_counts)
        fig, ax = plt.subplots(figsize=(11, 9))
        ax.axis("off")
        ax.set_title(f"Repository: {repo_name} \u2014 Rank-ordered top 20 classes", fontsize=13, pad=20)
        rows = []
        for rank, (code, cnt) in enumerate(top20, start=1):
            pct = 100 * cnt / total if total else 0
            rows.append([str(rank), full_class_name(code, div_titles), str(cnt), f"{pct:.1f}%"])
        table = ax.table(
            cellText=rows,
            colLabels=["Rank", "ISIC class", "Count", "% of classified"],
            colWidths=[0.06, 0.62, 0.12, 0.18],
            loc="center",
            cellLoc="left",
        )
        table.auto_set_font_size(False)
        table.set_fontsize(8)
        table.scale(1, 1.5)
        for (r, c), cell in table.get_celld().items():
            if r == 0:
                cell.set_facecolor("#4472C4")
                cell.set_text_props(color="white", fontweight="bold")
        pdf.savefig(fig)
        plt.close(fig)

    def comments_page(pdf, repo_name, comments):
        fig = plt.figure(figsize=(11, 8.5))
        fig.text(0.08, 0.92, f"Repository: {repo_name} \u2014 Comments", fontsize=14, fontweight="bold")
        wrapped = "\n\n".join(textwrap.fill(c, 100) for c in comments)
        fig.text(0.08, 0.85, wrapped, fontsize=10, va="top")
        plt.axis("off")
        pdf.savefig(fig)
        plt.close(fig)

    with PdfPages(out_path) as pdf:
        title_page(pdf)
        for repo_id, repo_name in repos:
            type_counts, class_counts = get_repo_stats(conn, repo_id)
            type_distribution_page(pdf, repo_name, type_counts)
            histogram_page(pdf, repo_name, class_counts)
            table_page(pdf, repo_name, class_counts)
            comments_page(pdf, repo_name, comments_by_repo.get(repo_id, ["(no comments provided)"]))
    print(f"[pdf]   report -> {out_path}")


# ===========================================================================
# Default per-repository comments (edit these to match your own findings)
# ===========================================================================

DEFAULT_COMMENTS = {
    5: [
        "Education (Q85) is by far the most common primary class for DANS projects, "
        "consistent with the repository's strong base of interview and life-history "
        "studies conducted in academic/educational settings.",
        "Several high-ranking classes (H49 Land transport, H52 Warehousing, B09 Mining "
        "support) look implausible for qualitative interview data at first glance. "
        "Because classification here is TF-IDF text matching between project "
        "title/description/keywords and the ISIC division reference text, short or "
        "generic titles with little descriptive text are prone to matching on "
        "incidental vocabulary overlap rather than true subject matter - a known "
        "limitation of a metadata-only classifier, flagged here as a data-quality "
        "issue rather than smoothed over.",
        "15 projects were classified as QDA_PROJECT (contain an actual REFI-QDA / "
        "MaxQDA / ATLAS.ti analysis file); the remaining classified projects are "
        "QD_PROJECT (primary data files present, but no analysis-software file).",
    ],
    16: [
        "No histogram or top-20 table is shown for this repository: nearly every "
        "project was classified NOT_A_PROJECT, so none were run through the ISIC "
        "classifier (per the assignment, only QDA_PROJECT and QD_PROJECT are "
        "classified).",
        "Root cause: the FILES table has very few rows for this repository "
        "compared to the other one, which traces back to Part 1's scraper not "
        "successfully enumerating each project's files - a data-acquisition gap "
        "worth revisiting before final submission, since it silently excludes an "
        "entire repository from classification.",
    ],
}


# ===========================================================================
# Main
# ===========================================================================

def main():
    ap = argparse.ArgumentParser(description="SQ26 Part 2 - classification pipeline (single file)")
    ap.add_argument("--db", required=True, help="Path to the Part 1 SQLite database")
    ap.add_argument("--isic-xlsx", required=True, help="Path to ISIC5_Exp_Notes_11Mar2024.xlsx")
    ap.add_argument("--out-prefix", default="sq26-classification",
                     help="Prefix for output files (default: sq26-classification)")
    args = ap.parse_args()

    out_db = f"{args.out_prefix}.db"
    out_xlsx = f"{args.out_prefix}-results.xlsx"
    out_pdf = f"{args.out_prefix}-report.pdf"

    if os.path.abspath(out_db) != os.path.abspath(args.db):
        shutil.copyfile(args.db, out_db)
        print(f"[setup] copied {args.db} -> {out_db} (original left untouched)")

    conn = sqlite3.connect(out_db)

    migrate_schema(conn)
    classify_project_types(conn)

    divisions = extract_isic_divisions(args.isic_xlsx)
    classify_isic(conn, divisions)

    export_xlsx(conn, out_xlsx)

    cur = conn.cursor()
    cur.execute("SELECT id, name FROM REPOSITORIES ORDER BY id")
    repos = cur.fetchall()
    generate_report(conn, divisions, repos, DEFAULT_COMMENTS, out_pdf)

    conn.close()
    print("\n[done] Part 2 pipeline complete.")
    print(f"       DB   : {out_db}")
    print(f"       XLSX : {out_xlsx}")
    print(f"       PDF  : {out_pdf}")


if __name__ == "__main__":
    main()
